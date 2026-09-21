from __future__ import annotations

from io import BytesIO
from typing import Any, Iterable, Sequence

from openpyxl import load_workbook


# --------------------------------------------------
# Generel tekstvalidering
# --------------------------------------------------
def normalize_required_text(
    *,
    name: str,
    value: str,
) -> str:
    """Validerer og normaliserer en obligatorisk tekstværdi."""
    normalized_name = _normalize_parameter_name(name)

    if not isinstance(value, str):
        raise TypeError(
            f"{normalized_name} skal være tekst. "
            f"Modtog: {type(value).__name__}."
        )

    normalized_value = value.strip()

    if not normalized_value:
        raise ValueError(
            f"{normalized_name} må ikke være tom."
        )

    return normalized_value


# --------------------------------------------------
# Generel id-validering
# --------------------------------------------------
def normalize_positive_id(
    *,
    name: str,
    value: int | str,
) -> int:
    """Validerer og normaliserer et positivt heltals-id."""
    normalized_name = _normalize_parameter_name(name)

    if isinstance(value, bool):
        raise TypeError(
            f"{normalized_name} må ikke være boolsk."
        )

    if isinstance(value, int):
        normalized_value = value
    elif isinstance(value, str):
        text_value = value.strip()

        if not text_value:
            raise ValueError(
                f"{normalized_name} må ikke være tom."
            )

        if not text_value.isdigit():
            raise ValueError(
                f"{normalized_name} skal være numerisk. "
                f"Modtog: {value!r}."
            )

        normalized_value = int(text_value)
    else:
        raise TypeError(
            f"{normalized_name} skal være et heltal eller "
            "en numerisk tekstværdi. "
            f"Modtog: {type(value).__name__}."
        )

    if normalized_value <= 0:
        raise ValueError(
            f"{normalized_name} skal være større end 0."
        )

    return normalized_value


def normalize_non_negative_int(
    *,
    name: str,
    value: int,
) -> int:
    """Validerer et heltal, som ikke må være negativt."""
    normalized_name = _normalize_parameter_name(name)

    if isinstance(value, bool) or not isinstance(value, int):
        raise TypeError(
            f"{normalized_name} skal være et heltal."
        )

    if value < 0:
        raise ValueError(
            f"{normalized_name} må ikke være negativ."
        )

    return value


# --------------------------------------------------
# Bool-validering
# --------------------------------------------------
def validate_bool(
    *,
    name: str,
    value: bool,
) -> bool:
    """Validerer og returnerer en boolsk værdi."""
    normalized_name = _normalize_parameter_name(name)

    if not isinstance(value, bool):
        raise TypeError(
            f"{normalized_name} skal være True eller False."
        )

    return value


# --------------------------------------------------
# Årsintervaller
# --------------------------------------------------
def validate_year_range(
    *,
    from_name: str,
    from_year: int,
    to_name: str,
    to_year: int,
) -> tuple[int, int]:
    """Validerer et interval mellem to heltalsår."""
    normalized_from_name = _normalize_parameter_name(
        from_name
    )
    normalized_to_name = _normalize_parameter_name(
        to_name
    )

    if isinstance(from_year, bool) or not isinstance(
        from_year,
        int,
    ):
        raise TypeError(
            f"{normalized_from_name} skal være et heltal."
        )

    if isinstance(to_year, bool) or not isinstance(
        to_year,
        int,
    ):
        raise TypeError(
            f"{normalized_to_name} skal være et heltal."
        )

    if from_year < 0:
        raise ValueError(
            f"{normalized_from_name} må ikke være negativ."
        )

    if to_year < 0:
        raise ValueError(
            f"{normalized_to_name} må ikke være negativ."
        )

    if from_year > to_year:
        raise ValueError(
            f"{normalized_from_name} må ikke være større "
            f"end {normalized_to_name}."
        )

    return from_year, to_year


# --------------------------------------------------
# Eksportkolonner
# --------------------------------------------------
def normalize_columns(
    *,
    columns: Sequence[str] | None,
    default_columns: Sequence[str],
) -> list[str]:
    """Validerer kolonner til en Insubiz-eksport."""
    if isinstance(default_columns, (str, bytes)):
        raise TypeError(
            "default_columns skal være en sekvens af tekst."
        )

    selected_columns: Sequence[str] = (
        default_columns
        if columns is None
        else columns
    )

    if isinstance(selected_columns, (str, bytes)):
        raise TypeError(
            "columns skal være en sekvens af tekst."
        )

    if not selected_columns:
        raise ValueError(
            "columns må ikke være tom."
        )

    normalized_columns: list[str] = []

    for column_number, column in enumerate(
        selected_columns,
        start=1,
    ):
        if not isinstance(column, str):
            raise TypeError(
                "Alle værdier i columns skal være tekst. "
                f"Kolonneposition: {column_number}. "
                f"Modtog: {type(column).__name__}."
            )

        normalized_column = column.strip()

        if not normalized_column:
            raise ValueError(
                "columns må ikke indeholde tomme værdier. "
                f"Kolonneposition: {column_number}."
            )

        normalized_columns.append(normalized_column)

    return normalized_columns


# --------------------------------------------------
# Binær response til tekst
# --------------------------------------------------
def decode_response(
    *,
    content: bytes,
    max_length: int = 500,
) -> str:
    """Konverterer en binær fejlresponse til læsbar tekst."""
    if not isinstance(content, bytes):
        raise TypeError(
            "content skal være bytes."
        )

    if isinstance(max_length, bool) or not isinstance(
        max_length,
        int,
    ):
        raise TypeError(
            "max_length skal være et heltal."
        )

    if max_length <= 0:
        raise ValueError(
            "max_length skal være større end 0."
        )

    try:
        response_text = content.decode("utf-8")
    except UnicodeDecodeError:
        return repr(content[:max_length])

    return response_text[:max_length]


# --------------------------------------------------
# Excel-response
# --------------------------------------------------
def parse_excel_response(
    *,
    content: bytes,
    content_type: str,
    resource_name: str,
) -> list[dict[str, Any]]:
    """
    Konverterer en binær Excel-response til dictionaries.

    Funktionen bruges af skade-, køretøjs- og
    ejendomseksporterne.
    """
    normalized_resource_name = normalize_required_text(
        name="resource_name",
        value=resource_name,
    )

    if not isinstance(content, bytes):
        raise RuntimeError(
            f"{normalized_resource_name} havde et uventet "
            "format. Forventede bytes, men modtog "
            f"{type(content).__name__}."
        )

    if not content:
        raise RuntimeError(
            f"{normalized_resource_name} var tom."
        )

    if not isinstance(content_type, str):
        raise TypeError(
            "content_type skal være tekst."
        )

    _raise_for_non_excel_response(
        content=content,
        content_type=content_type,
        resource_name=normalized_resource_name,
    )

    try:
        workbook = load_workbook(
            filename=BytesIO(content),
            read_only=True,
            data_only=True,
        )
    except Exception as error:
        raise RuntimeError(
            f"{normalized_resource_name} kunne ikke læses "
            "som en Excel-fil. "
            f"Content-Type: {content_type!r}."
        ) from error

    try:
        worksheet = workbook.active

        if worksheet is None:
            raise RuntimeError(
                f"{normalized_resource_name} indeholder "
                "ikke et aktivt regneark."
            )

        rows = worksheet.iter_rows(
            values_only=True,
        )
        header_row = next(rows, None)

        if header_row is None:
            return []

        headers = create_unique_headers(
            header_row=header_row,
        )

        result: list[dict[str, Any]] = []

        for row in rows:
            if row_is_empty(row=row):
                continue

            row_values = list(row)

            if len(row_values) < len(headers):
                row_values.extend(
                    [None]
                    * (
                        len(headers)
                        - len(row_values)
                    )
                )

            result.append(
                dict(
                    zip(
                        headers,
                        row_values[:len(headers)],
                    )
                )
            )

        return result
    finally:
        workbook.close()


def create_unique_headers(
    *,
    header_row: Iterable[Any],
) -> list[str]:
    """Opretter stabile og unikke Excel-kolonnenavne."""
    if isinstance(header_row, (str, bytes)):
        raise TypeError(
            "header_row skal være en iterable af "
            "celleværdier."
        )

    headers: list[str] = []
    used_headers: set[str] = set()

    for column_number, value in enumerate(
        header_row,
        start=1,
    ):
        base_header = _create_base_header(
            value=value,
            column_number=column_number,
        )

        header = base_header
        duplicate_number = 2

        while header in used_headers:
            header = (
                f"{base_header}_{duplicate_number}"
            )
            duplicate_number += 1

        used_headers.add(header)
        headers.append(header)

    return headers


def row_is_empty(
    *,
    row: Iterable[Any],
) -> bool:
    """Returnerer True, hvis hele Excel-rækken er tom."""
    if isinstance(row, (str, bytes)):
        raise TypeError(
            "row skal være en iterable af celleværdier."
        )

    return all(
        value is None
        or (
            isinstance(value, str)
            and not value.strip()
        )
        for value in row
    )


# --------------------------------------------------
# Interne hjælpefunktioner
# --------------------------------------------------
def _normalize_parameter_name(name: str) -> str:
    """Validerer et parameternavn til fejlbeskeder."""
    if not isinstance(name, str):
        raise TypeError(
            "name skal være tekst."
        )

    normalized_name = name.strip()

    if not normalized_name:
        raise ValueError(
            "name må ikke være tom."
        )

    return normalized_name


def _raise_for_non_excel_response(
    *,
    content: bytes,
    content_type: str,
    resource_name: str,
) -> None:
    """Stopper ved JSON- eller HTML-svar i stedet for Excel."""
    normalized_content_type = content_type.strip().casefold()

    if "json" in normalized_content_type:
        response_text = decode_response(
            content=content,
        )
        raise RuntimeError(
            "Insubiz returnerede JSON i stedet for "
            f"{resource_name}. "
            f"Response: {response_text}"
        )

    if "html" in normalized_content_type:
        response_text = decode_response(
            content=content,
        )
        raise RuntimeError(
            "Insubiz returnerede HTML i stedet for "
            f"{resource_name}. "
            "Login-sessionen kan være udløbet. "
            f"Response: {response_text}"
        )


def _create_base_header(
    *,
    value: Any,
    column_number: int,
) -> str:
    """Opretter grundnavnet til en Excel-kolonne."""
    if isinstance(column_number, bool) or not isinstance(
        column_number,
        int,
    ):
        raise TypeError(
            "column_number skal være et heltal."
        )

    if column_number <= 0:
        raise ValueError(
            "column_number skal være større end 0."
        )

    if value is None:
        return f"column_{column_number}"

    normalized_value = str(value).strip()

    if not normalized_value:
        return f"column_{column_number}"

    return normalized_value


__all__ = [
    "create_unique_headers",
    "decode_response",
    "normalize_columns",
    "normalize_non_negative_int",
    "normalize_positive_id",
    "normalize_required_text",
    "parse_excel_response",
    "row_is_empty",
    "validate_bool",
    "validate_year_range",
]
