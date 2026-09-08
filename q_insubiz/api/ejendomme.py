from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from q_insubiz.api.client import InsubizApiClient


# --------------------------------------------------
# Endpoints
# --------------------------------------------------
EJENDOMME_LISTE_ENDPOINT = (
    "/ImportExport/ExportLocations"
)


# --------------------------------------------------
# Query-parametre
# --------------------------------------------------
EJENDOMME_LISTE_PARAMS = {
    "customerId": 0,
    "showTreeData": "false",
    "activeOnly": "false",
}


# --------------------------------------------------
# Standardkolonner
# --------------------------------------------------
# Erstat værdierne her, hvis Network-fanen viser
# andre præcise kolonnenavne for ExportLocations.
EJENDOMME_LISTE_COLUMNS = [
    "Name",
    "Address",
]


async def EJENDOMME_LISTE(
    api_client: InsubizApiClient,
    *,
    customer_id: int = 0,
    show_tree_data: bool = False,
    active_only: bool = False,
    columns: list[str] | None = None,
) -> list[dict[str, Any]]:
    """
    Henter en liste over ejendomme fra Insubiz.

    Eksporten behandles direkte i hukommelsen.
    Der bliver ikke gemt en fil på disken.

    Output:
    En liste af dictionaries, hvor hver dictionary
    repræsenterer en række fra Excel-eksporten.
    """

    if (
        not isinstance(customer_id, int)
        or isinstance(customer_id, bool)
    ):
        raise ValueError(
            "customer_id skal være et heltal."
        )

    if customer_id < 0:
        raise ValueError(
            "customer_id må ikke være negativ."
        )

    if not isinstance(show_tree_data, bool):
        raise ValueError(
            "show_tree_data skal være True eller False."
        )

    if not isinstance(active_only, bool):
        raise ValueError(
            "active_only skal være True eller False."
        )

    selected_columns = (
        list(EJENDOMME_LISTE_COLUMNS)
        if columns is None
        else columns
    )

    if not isinstance(selected_columns, list):
        raise ValueError(
            "columns skal være en liste."
        )

    if not selected_columns:
        raise ValueError(
            "columns må ikke være tom."
        )

    normalized_columns: list[str] = []

    for column in selected_columns:
        if not isinstance(column, str):
            raise ValueError(
                "Alle værdier i columns skal være tekst."
            )

        normalized_column = column.strip()

        if not normalized_column:
            raise ValueError(
                "columns må ikke indeholde tomme værdier."
            )

        normalized_columns.append(normalized_column)

    content, content_type = await api_client.download(
        endpoint=EJENDOMME_LISTE_ENDPOINT,
        method="POST",
        params={
            "customerId": customer_id,
            "showTreeData": str(show_tree_data).lower(),
            "activeOnly": str(active_only).lower(),
        },
        json_body=normalized_columns,
    )

    if not isinstance(content, bytes):
        raise RuntimeError(
            "Insubiz-eksporten havde et uventet format. "
            "Forventede bytes, "
            f"men modtog {type(content).__name__}."
        )

    if not content:
        raise RuntimeError(
            "Insubiz-eksporten var tom."
        )

    if "application/json" in content_type.lower():
        try:
            response_text = content.decode("utf-8")
        except UnicodeDecodeError:
            response_text = repr(content[:500])

        raise RuntimeError(
            "Insubiz returnerede JSON i stedet for "
            "en Excel-eksport. "
            f"Response: {response_text[:500]}"
        )

    try:
        workbook = load_workbook(
            filename=BytesIO(content),
            read_only=True,
            data_only=True,
        )
    except Exception as error:
        raise RuntimeError(
            "Insubiz-responsen kunne ikke læses "
            "som en Excel-fil. "
            f"Content-Type: {content_type!r}."
        ) from error

    try:
        worksheet = workbook.active

        if worksheet is None:
            raise RuntimeError(
                "Excel-eksporten indeholder ikke "
                "et aktivt regneark."
            )

        row_iterator = worksheet.iter_rows(
            values_only=True
        )

        header_row = next(row_iterator, None)

        if header_row is None:
            raise RuntimeError(
                "Excel-eksporten indeholder ingen rækker."
            )

        headers: list[str] = []

        for column_number, header_value in enumerate(
            header_row,
            start=1,
        ):
            if (
                header_value is not None
                and str(header_value).strip()
            ):
                header_name = str(header_value).strip()
            else:
                header_name = f"column_{column_number}"

            headers.append(header_name)

        results: list[dict[str, Any]] = []

        for row in row_iterator:
            row_is_empty = all(
                value is None
                or str(value).strip() == ""
                for value in row
            )

            if row_is_empty:
                continue

            row_values = list(row)

            if len(row_values) < len(headers):
                row_values.extend(
                    [None] * (len(headers) - len(row_values))
                )

            row_values = row_values[:len(headers)]

            row_result = dict(
                zip(headers, row_values)
            )

            results.append(row_result)

        return results

    finally:
        workbook.close()
